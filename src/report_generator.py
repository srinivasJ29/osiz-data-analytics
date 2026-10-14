import io
import os
import pandas as pd
from typing import Dict, Any, List, Optional
from datetime import datetime

# ReportLab imports
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, Image
)
from reportlab.pdfgen import canvas

# python-docx imports
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

class NumberedCanvas(canvas.Canvas):
    """Adds 'Page X of Y' to every page of the ReportLab PDF."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 9)
        self.setFillColor(colors.HexColor("#64748b"))
        
        # Header (pages after title page)
        if self._pageNumber > 1:
            self.drawString(54, 750, "Smart Data Analyzer — Executive Analytical Report")
            self.setStrokeColor(colors.HexColor("#cbd5e1"))
            self.setLineWidth(0.5)
            self.line(54, 744, 558, 744)

        # Footer
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 40, page_str)
        self.drawString(54, 40, f"Confidential & Automated • Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}")
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.5)
        self.line(54, 50, 558, 50)
        self.restoreState()

def generate_pdf_report(
    metadata: Dict[str, Any],
    stats_df: pd.DataFrame,
    insights: List[Dict[str, str]],
    audit_history: List[Dict[str, Any]],
    scope_name: str = "Cleaned",
    report_title: str = "Analytical Data Report",
    author_name: str = "Data Analytics Team",
    selected_sections: Optional[List[str]] = None
) -> bytes:
    """
    Generates a PDF analytical report using ReportLab with clean typography and styled tables.
    """
    if selected_sections is None:
        selected_sections = ["Dataset Overview", "Data Quality & Operations", "Statistical Summary", "Automated Insights"]

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    
    # Custom Typography
    title_style = ParagraphStyle(
        'DocTitle', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=24, leading=28,
        textColor=colors.HexColor('#0f172a'), spaceAfter=8
    )
    subtitle_style = ParagraphStyle(
        'DocSubTitle', parent=styles['Normal'],
        fontName='Helvetica', fontSize=12, leading=16,
        textColor=colors.HexColor('#475569'), spaceAfter=18
    )
    heading_style = ParagraphStyle(
        'SectionHeading', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=15, leading=19,
        textColor=colors.HexColor('#1e3a8a'), spaceBefore=14, spaceAfter=8
    )
    body_style = ParagraphStyle(
        'Body', parent=styles['Normal'],
        fontName='Helvetica', fontSize=10, leading=14,
        textColor=colors.HexColor('#334155'), spaceAfter=6
    )
    table_cell = ParagraphStyle(
        'TableCell', parent=styles['Normal'],
        fontName='Helvetica', fontSize=8, leading=11,
        textColor=colors.HexColor('#1e293b')
    )
    table_hdr = ParagraphStyle(
        'TableHdr', parent=styles['Normal'],
        fontName='Helvetica-Bold', fontSize=8, leading=11,
        textColor=colors.white
    )

    story = []

    # Title & Header
    story.append(Paragraph(report_title, title_style))
    story.append(Paragraph(f"Author / Analyst: {author_name} | Scope: {scope_name} Data | Generated: {datetime.now().strftime('%B %d, %Y')}", subtitle_style))
    story.append(Spacer(1, 10))

    # SECTION 1: Dataset Overview
    if "Dataset Overview" in selected_sections:
        story.append(Paragraph("1. Executive Dataset Overview", heading_style))
        story.append(Paragraph("The key architectural dimensions and profiling metadata of the examined dataset are outlined below:", body_style))
        
        overview_data = [
            [Paragraph("<b>Metric</b>", table_hdr), Paragraph("<b>Value</b>", table_hdr)],
            [Paragraph("Source Filename", table_cell), Paragraph(str(metadata.get("filename", "N/A")), table_cell)],
            [Paragraph("Total Records (Rows)", table_cell), Paragraph(f"{metadata.get('total_rows', 0):,}", table_cell)],
            [Paragraph("Total Attributes (Columns)", table_cell), Paragraph(str(metadata.get("total_cols", 0)), table_cell)],
            [Paragraph("Numerical Columns", table_cell), Paragraph(str(len(metadata.get("numerical_columns", []))), table_cell)],
            [Paragraph("Categorical Columns", table_cell), Paragraph(str(len(metadata.get("categorical_columns", []))), table_cell)],
            [Paragraph("Identified Duplicate Records", table_cell), Paragraph(str(metadata.get("duplicate_rows", 0)), table_cell)]
        ]
        
        t_overview = Table(overview_data, colWidths=[200, 304])
        t_overview.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a8a')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ]))
        story.append(t_overview)
        story.append(Spacer(1, 14))

    # SECTION 2: Data Quality & Operations Log
    if "Data Quality & Operations" in selected_sections:
        story.append(Paragraph("2. Data Quality & Transformations Audit", heading_style))
        if audit_history:
            story.append(Paragraph("Chronological log of verified cleaning and feature adjustments applied to this dataset:", body_style))
            audit_rows = [[Paragraph("<b>#</b>", table_hdr), Paragraph("<b>Action Performed</b>", table_hdr), Paragraph("<b>Affected Records/Cols</b>", table_hdr), Paragraph("<b>Timestamp</b>", table_hdr)]]
            for i, item in enumerate(audit_history, 1):
                audit_rows.append([
                    Paragraph(str(i), table_cell),
                    Paragraph(str(item.get("action", "")), table_cell),
                    Paragraph(str(item.get("affected", 0)), table_cell),
                    Paragraph(str(item.get("timestamp", "")), table_cell)
                ])
            t_audit = Table(audit_rows, colWidths=[30, 260, 114, 100])
            t_audit.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a8a')),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
                ('TOPPADDING', (0, 0), (-1, -1), 4),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
            ]))
            story.append(t_audit)
        else:
            story.append(Paragraph("<i>No data cleaning or modifying transformations have been applied (baseline original dataset).</i>", body_style))
        story.append(Spacer(1, 14))

    # SECTION 3: Statistical Summary
    if "Statistical Summary" in selected_sections and stats_df is not None and not stats_df.empty:
        story.append(Paragraph("3. Numerical Statistical Profile", heading_style))
        story.append(Paragraph("Summary parametric metrics computed for numerical features:", body_style))
        
        stat_cols = ["Column", "Count", "Mean", "Std Dev", "Min", "Median", "Max"]
        headers = [Paragraph(f"<b>{c}</b>", table_hdr) for c in stat_cols]
        stat_rows = [headers]
        
        # Take top 10 columns for readability in table
        for _, r in stats_df.head(10).iterrows():
            row_items = [
                Paragraph(str(r.get("Column", "")), table_cell),
                Paragraph(str(r.get("Count", "")), table_cell),
                Paragraph(str(r.get("Mean", "")), table_cell),
                Paragraph(str(r.get("Std Dev", "")), table_cell),
                Paragraph(str(r.get("Min", "")), table_cell),
                Paragraph(str(r.get("Median (50%)", r.get("Median", ""))), table_cell),
                Paragraph(str(r.get("Max", "")), table_cell),
            ]
            stat_rows.append(row_items)
            
        t_stats = Table(stat_rows, colWidths=[114, 55, 65, 65, 65, 65, 75])
        t_stats.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1e3a8a')),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ]))
        story.append(t_stats)
        story.append(Spacer(1, 14))

    # SECTION 4: Automated Insights
    if "Automated Insights" in selected_sections and insights:
        story.append(Paragraph("4. Automated Analytical Findings & Insights", heading_style))
        for item in insights:
            title_p = Paragraph(f"• <b>{item.get('title', 'Insight')}</b>", ParagraphStyle('BulletHdr', parent=body_style, fontName='Helvetica-Bold', textColor=colors.HexColor('#1e3a8a')))
            text_p = Paragraph(f"{item.get('text', '')}", ParagraphStyle('BulletText', parent=body_style, leftIndent=12))
            story.append(title_p)
            story.append(text_p)
            story.append(Spacer(1, 4))
        story.append(Spacer(1, 10))

    # Build document
    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()

def generate_word_report(
    metadata: Dict[str, Any],
    stats_df: pd.DataFrame,
    insights: List[Dict[str, str]],
    audit_history: List[Dict[str, Any]],
    scope_name: str = "Cleaned",
    report_title: str = "Analytical Data Report",
    author_name: str = "Data Analytics Team"
) -> bytes:
    """
    Generates a Word document (.docx) analytical report using python-docx.
    """
    doc = Document()
    
    # Title
    title_p = doc.add_heading(report_title, level=0)
    title_p.paragraph_format.space_after = Pt(4)
    
    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Author / Analyst: {author_name} | Scope: {scope_name} | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}").italic = True
    meta_p.paragraph_format.space_after = Pt(16)

    # 1. Overview
    doc.add_heading("1. Executive Dataset Overview", level=1)
    doc.add_paragraph(f"Analysis conducted on dataset: {metadata.get('filename', 'N/A')}")
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Light Shading'
    hdr = table.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Value"
    
    metrics = [
        ("Total Records (Rows)", f"{metadata.get('total_rows', 0):,}"),
        ("Total Attributes (Columns)", str(metadata.get('total_cols', 0))),
        ("Numerical Columns", str(len(metadata.get('numerical_columns', [])))),
        ("Categorical Columns", str(len(metadata.get('categorical_columns', [])))),
        ("Duplicate Rows", str(metadata.get('duplicate_rows', 0)))
    ]
    for m, val in metrics:
        row = table.add_row().cells
        row[0].text = m
        row[1].text = val

    # 2. Quality & Operations
    doc.add_heading("2. Data Quality & Transformations Audit", level=1)
    if audit_history:
        a_table = doc.add_table(rows=1, cols=3)
        a_table.style = 'Light Shading'
        a_hdr = a_table.rows[0].cells
        a_hdr[0].text = "Action"
        a_hdr[1].text = "Affected Count"
        a_hdr[2].text = "Timestamp"
        for item in audit_history:
            row = a_table.add_row().cells
            row[0].text = str(item.get("action", ""))
            row[1].text = str(item.get("affected", 0))
            row[2].text = str(item.get("timestamp", ""))
    else:
        doc.add_paragraph("No cleaning transformations applied to this dataset.")

    # 3. Numerical Summary
    if stats_df is not None and not stats_df.empty:
        doc.add_heading("3. Numerical Statistical Profile", level=1)
        stat_cols = ["Column", "Count", "Mean", "Std Dev", "Min", "Median", "Max"]
        s_table = doc.add_table(rows=1, cols=len(stat_cols))
        s_table.style = 'Light Shading'
        for i, col in enumerate(stat_cols):
            s_table.rows[0].cells[i].text = col
        for _, r in stats_df.head(15).iterrows():
            row = s_table.add_row().cells
            row[0].text = str(r.get("Column", ""))
            row[1].text = str(r.get("Count", ""))
            row[2].text = str(r.get("Mean", ""))
            row[3].text = str(r.get("Std Dev", ""))
            row[4].text = str(r.get("Min", ""))
            row[5].text = str(r.get("Median (50%)", r.get("Median", "")))
            row[6].text = str(r.get("Max", ""))

    # 4. Insights
    if insights:
        doc.add_heading("4. Key Analytical Insights", level=1)
        for item in insights:
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(f"{item.get('title', 'Insight')}: ")
            run.bold = True
            p.add_run(item.get('text', ''))

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def export_to_csv(df: pd.DataFrame) -> bytes:
    """Exports DataFrame to UTF-8 CSV bytes."""
    return df.to_csv(index=False).encode('utf-8')

def export_to_excel(df: pd.DataFrame, sheet_name: str = "Data") -> bytes:
    """Exports DataFrame to Excel workbook bytes via openpyxl."""
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=sheet_name)
    buffer.seek(0)
    return buffer.getvalue()
